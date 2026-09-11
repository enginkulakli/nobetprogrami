import io
from datetime import date
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse

from deps import db, current_station
from helpers import get_active_personnel_list, personnel_card_data, MONTHS_TR

router = APIRouter(prefix="/api", tags=["exports"])
ROOT_DIR = Path(__file__).parent.parent


def _register_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts_dir = ROOT_DIR / "fonts"
    try:
        if "LiberationSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("LiberationSans", str(fonts_dir / "LiberationSans-Regular.ttf")))
            pdfmetrics.registerFont(TTFont("LiberationSans-Bold", str(fonts_dir / "LiberationSans-Bold.ttf")))
        return "LiberationSans"
    except Exception:
        return "Helvetica"


def _card_elements(p: dict, yil: int, ay: int, items: list, toplam: int) -> list:
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import Table, TableStyle, Paragraph, Spacer

    font = _register_pdf_font()
    bold = f"{font}-Bold" if font != "Helvetica" else "Helvetica-Bold"
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontName=bold)
    body = ParagraphStyle("b", parent=styles["Normal"], fontName=font, fontSize=11)
    elements = [
        Paragraph(f"112 Nöbet Kartı — {MONTHS_TR[ay - 1]} {yil}", title),
        Paragraph(f"<b>{p['ad_soyad']}</b> · {len(items)} nöbet · {toplam} saat", body),
        Spacer(1, 12),
    ]
    data = [["Tarih", "Gün", "Görev", "Saat", "Not"]]
    for i in items:
        data.append([i["tarih"], i["gun"], i["gorev"], i["saat_araligi"],
                     i["tatil"] or ("Hafta sonu" if i["hafta_sonu"] else "")])
    t = Table(data, repeatRows=1, colWidths=[80, 45, 120, 80, 180])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    elements.append(t)
    return elements


def _build_pdf(elements: list) -> io.BytesIO:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(elements)
    buf.seek(0)
    return buf


@router.get("/schedule/{yil}/{ay}/personnel/{pid}/export/pdf")
async def personnel_card_pdf(yil: int, ay: int, pid: str, istasyon_id: str = Depends(current_station)):
    p, items, toplam = await personnel_card_data(yil, ay, pid, istasyon_id)
    buf = _build_pdf(_card_elements(p, yil, ay, items, toplam))
    fname = quote(f"nobet_karti_{p['ad_soyad'].replace(' ', '_')}_{yil}_{ay:02d}.pdf")
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fname}"})


async def _all_cards(yil: int, ay: int, istasyon_id: str):
    doc = await db.schedules.find_one({"yil": yil, "ay": ay, "istasyon_id": istasyon_id}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Liste yok")
    assigned = {a["personel_id"] for a in doc.get("assignments", [])}
    personnel = sorted([p for p in await get_active_personnel_list(istasyon_id) if p["id"] in assigned],
                       key=lambda p: p["ad_soyad"])
    out = []
    for p in personnel:
        _, items, toplam = await personnel_card_data(yil, ay, p["id"], istasyon_id)
        out.append((p, items, toplam))
    return out


@router.get("/schedule/{yil}/{ay}/personnel-cards/export/pdf")
async def all_cards_pdf(yil: int, ay: int, istasyon_id: str = Depends(current_station)):
    from reportlab.platypus import PageBreak
    cards = await _all_cards(yil, ay, istasyon_id)
    elements = []
    for idx, (p, items, toplam) in enumerate(cards):
        if idx:
            elements.append(PageBreak())
        elements += _card_elements(p, yil, ay, items, toplam)
    return StreamingResponse(_build_pdf(elements), media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=nobet_kartlari_{yil}_{ay:02d}.pdf"})


@router.get("/schedule/{yil}/{ay}/personnel-cards/export/zip")
async def all_cards_zip(yil: int, ay: int, istasyon_id: str = Depends(current_station)):
    import zipfile
    cards = await _all_cards(yil, ay, istasyon_id)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for p, items, toplam in cards:
            pdf = _build_pdf(_card_elements(p, yil, ay, items, toplam))
            zf.writestr(f"nobet_karti_{p['ad_soyad'].replace(' ', '_')}_{yil}_{ay:02d}.pdf", pdf.getvalue())
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": f"attachment; filename=nobet_kartlari_{yil}_{ay:02d}.zip"})


@router.get("/schedule/{yil}/{ay}/export/excel")
async def export_excel(yil: int, ay: int, istasyon_id: str = Depends(current_station)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    doc = await db.schedules.find_one({"yil": yil, "ay": ay, "istasyon_id": istasyon_id}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Liste yok")
    personnel = await get_active_personnel_list(istasyon_id)
    pmap = {p["id"]: p["ad_soyad"] for p in personnel}

    wb = Workbook()
    ws = wb.active
    ws.title = f"{yil}-{ay:02d}"
    headers = ["Tarih", "Gün", "Sağlıkçı 1", "Sağlıkçı 2", "Sürücü 09-21", "Sürücü 21-09", "24s Sürücü"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2A4A")
        c.alignment = Alignment(horizontal="center")

    gunler_tr = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cts", "Paz"]
    for day in doc["gunler"]:
        d = date.fromisoformat(day["tarih"])
        ws.append([
            day["tarih"], gunler_tr[d.weekday()],
            pmap.get(day.get("saglikci_1") or "", ""),
            pmap.get(day.get("saglikci_2") or "", ""),
            pmap.get(day.get("surucu_gunduz") or "", ""),
            pmap.get(day.get("surucu_gece") or "", ""),
            pmap.get(day.get("surucu_24") or "", ""),
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=nobet_{yil}_{ay:02d}.xlsx"},
    )


@router.get("/schedule/{yil}/{ay}/export/pdf")
async def export_pdf(yil: int, ay: int, istasyon_id: str = Depends(current_station)):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    doc_db = await db.schedules.find_one({"yil": yil, "ay": ay, "istasyon_id": istasyon_id}, {"_id": 0})
    if not doc_db:
        raise HTTPException(404, "Liste yok")
    personnel = await get_active_personnel_list(istasyon_id)
    pmap = {p["id"]: p["ad_soyad"] for p in personnel}

    font = _register_pdf_font()
    bold = f"{font}-Bold" if font != "Helvetica" else "Helvetica-Bold"
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = bold
    elements = [Paragraph(f"<b>112 Nöbet Listesi - {yil}/{ay:02d}</b>", styles["Title"]), Spacer(1, 10)]

    gunler_tr = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cts", "Paz"]
    data = [["Tarih", "Gün", "Sağlıkçı 1", "Sağlıkçı 2", "Sürücü 09-21", "Sürücü 21-09", "24s Sürücü"]]
    for day in doc_db["gunler"]:
        d = date.fromisoformat(day["tarih"])
        data.append([
            day["tarih"], gunler_tr[d.weekday()],
            pmap.get(day.get("saglikci_1") or "", "-"),
            pmap.get(day.get("saglikci_2") or "", "-"),
            pmap.get(day.get("surucu_gunduz") or "", "-"),
            pmap.get(day.get("surucu_gece") or "", "-"),
            pmap.get(day.get("surucu_24") or "", "-"),
        ])

    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2A4A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    elements.append(t)
    pdf.build(elements)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=nobet_{yil}_{ay:02d}.pdf"},
    )
